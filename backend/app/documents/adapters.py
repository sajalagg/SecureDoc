import io
from pathlib import Path
from typing import Dict, Type

import docx
from pypdf import PdfReader

MAX_TEXT_FILE_BYTES = 1_000_000
MAX_DOCX_FILE_BYTES = 10_000_000
MAX_PDF_FILE_BYTES = 10_000_000


class DocumentFormatError(ValueError):
    """Raised when an uploaded file is not a safe, supported document."""


class TextDocumentAdapter:
    """Adapter for UTF-8 `.txt` files."""

    supported_extension = ".txt"

    @classmethod
    def read(cls, filename: str, content: bytes) -> str:
        """Validate and decode an uploaded UTF-8 text file without altering its contents."""
        if Path(filename).suffix.lower() != cls.supported_extension:
            raise DocumentFormatError("Only .txt files are supported in this milestone.")
        if len(content) > MAX_TEXT_FILE_BYTES:
            raise DocumentFormatError(f"Text files must be at most {MAX_TEXT_FILE_BYTES:,} bytes.")
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError as error:
            raise DocumentFormatError("Text files must use UTF-8 encoding.") from error

    @classmethod
    def write(cls, text: str) -> bytes:
        """Encode reconstructed normalized text as UTF-8 for a future download endpoint."""
        return text.encode("utf-8")


class DocxDocumentAdapter:
    """Adapter for Microsoft Word (.docx) documents.

    Extracts text from paragraphs and tables into normalized text for SecureDoc.
    """

    supported_extension = ".docx"

    @classmethod
    def read(cls, filename: str, content: bytes) -> str:
        """Validate and extract text from a .docx file."""
        if Path(filename).suffix.lower() != cls.supported_extension:
            raise DocumentFormatError("Only .docx files are supported by this adapter.")
        if len(content) > MAX_DOCX_FILE_BYTES:
            raise DocumentFormatError(f"DOCX files must be at most {MAX_DOCX_FILE_BYTES:,} bytes.")
        try:
            document = docx.Document(io.BytesIO(content))
        except Exception as error:
            raise DocumentFormatError(f"Invalid or corrupted DOCX file: {error}") from error

        lines = []
        for paragraph in document.paragraphs:
            if paragraph.text:
                lines.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    lines.append("\t".join(cells))

        text = "\n".join(lines).strip()
        if not text:
            raise DocumentFormatError("DOCX file contains no extractable text.")
        return text

    @classmethod
    def write(cls, text: str) -> bytes:
        """Encode normalized text into a new .docx document."""
        document = docx.Document()
        for line in text.splitlines():
            document.add_paragraph(line)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()


class PdfDocumentAdapter:
    """Adapter for PDF (.pdf) documents.

    Extracts text using pypdf.
    Limitations: Scanned or image-only PDFs contain no embedded text layer and
    require OCR preprocessing prior to upload. Encrypted/password-protected PDFs
    are rejected safely.
    """

    supported_extension = ".pdf"

    @classmethod
    def read(cls, filename: str, content: bytes) -> str:
        """Validate and extract text from an unencrypted PDF file."""
        if Path(filename).suffix.lower() != cls.supported_extension:
            raise DocumentFormatError("Only .pdf files are supported by this adapter.")
        if len(content) > MAX_PDF_FILE_BYTES:
            raise DocumentFormatError(f"PDF files must be at most {MAX_PDF_FILE_BYTES:,} bytes.")
        try:
            reader = PdfReader(io.BytesIO(content))
        except Exception as error:
            raise DocumentFormatError(f"Invalid or corrupted PDF file: {error}") from error

        if reader.is_encrypted:
            raise DocumentFormatError("Encrypted or password-protected PDF files are not supported.")

        pages = []
        try:
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    pages.append(extracted)
        except Exception as error:
            raise DocumentFormatError(f"Failed to extract text from PDF: {error}") from error

        text = "\n".join(pages).strip()
        if not text:
            raise DocumentFormatError(
                "PDF contains no extractable text. Scanned or image-only PDFs require OCR preprocessing, which is not supported."
            )
        return text


ADAPTER_REGISTRY: Dict[str, Type] = {
    ".txt": TextDocumentAdapter,
    ".docx": DocxDocumentAdapter,
    ".pdf": PdfDocumentAdapter,
}


def get_adapter_for_filename(filename: str):
    """Return the adapter class matching the file extension, or raise DocumentFormatError."""
    suffix = Path(filename).suffix.lower()
    adapter = ADAPTER_REGISTRY.get(suffix)
    if adapter is None:
        supported = ", ".join(sorted(ADAPTER_REGISTRY.keys()))
        raise DocumentFormatError(f"Unsupported file format '{suffix}'. Supported formats: {supported}.")
    return adapter


def read_document_file(filename: str, content: bytes) -> str:
    """Read and extract text from an uploaded file using the appropriate adapter."""
    adapter = get_adapter_for_filename(filename)
    return adapter.read(filename, content)
